"""Extract all paragraph and table text from a DOCX for requirements review."""

from __future__ import annotations

import argparse
import sys
from pathlib import Path

from docx import Document


def main() -> None:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    parser = argparse.ArgumentParser()
    parser.add_argument("docx", type=Path)
    parser.add_argument("--output", type=Path)
    args = parser.parse_args()

    document = Document(args.docx)
    lines: list[str] = ["PARAGRAPHS"]
    for index, paragraph in enumerate(document.paragraphs):
        lines.append(f"{index:03d} [{paragraph.style.name}] {paragraph.text}")

    lines.append("TABLES")
    for table_index, table in enumerate(document.tables):
        lines.append(f"TABLE {table_index}")
        for row in table.rows:
            cells = [cell.text.replace("\n", " / ") for cell in row.cells]
            lines.append(" | ".join(cells))

    output = "\n".join(lines)
    if args.output:
        args.output.parent.mkdir(parents=True, exist_ok=True)
        args.output.write_text(output, encoding="utf-8")
    print(output)


if __name__ == "__main__":
    main()
